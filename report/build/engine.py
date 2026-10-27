"""Document engine for the ClinXAI report: styles, sections, template pages, figures, tables, fields.

Formatting authority: 'TRAI - Micro_Project_Report_Format.docx' (Section 1 layout & typography).
Initial-page layout authority: 'TRAI - Micro-Project Report - Format for Initial Pages.docx' (copied verbatim,
then only the blanks are filled).
"""
from __future__ import annotations

import copy
import io
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

TEXT_WIDTH_IN = 6.0          # A4 8.27" - 1.25" left - 1.0" right = 6.02"
MAX_FIG_HEIGHT_IN = 8.0
BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
CODE_FILL = "F2F2F2"         # "White, Background 1, Darker 5%" = 5 % grey fill
HEADER_FILL = "D9E2F3"


# ------------------------------------------------------------------ low-level XML helpers
def set_rfonts(rpr, name: str) -> None:
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn("w:" + a), None)
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + a), name)


def set_color(rpr, hex_rgb: str) -> None:
    for c in rpr.findall(qn("w:color")):
        rpr.remove(c)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), hex_rgb)
    rpr.append(c)


def shade(element_pr, fill: str) -> None:
    for s in element_pr.findall(qn("w:shd")):
        element_pr.remove(s)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def para_borders(p, color="BFBFBF", sides=("top", "left", "bottom", "right"), sz="4", space="2"):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for side in sides:
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), space)
        e.set(qn("w:color"), color)
        pbdr.append(e)
    # pBdr must precede shd/tabs/spacing/ind/jc in CT_PPr; insert after keep/pageBreak elements
    anchor = None
    for tag in ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:jc", "w:rPr"):
        anchor = pPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(pbdr)
    else:
        pPr.append(pbdr)


def para_shading(p, fill=CODE_FILL):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    anchor = None
    for tag in ("w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:jc", "w:rPr"):
        anchor = pPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(shd)
    else:
        pPr.append(shd)


def add_field(paragraph, instr: str, placeholder: str = "", font_size=None):
    """Complex field (begin / instr / separate / placeholder / end) so Word can update it."""
    def r_with(child):
        r = OxmlElement("w:r")
        if font_size:
            rpr = OxmlElement("w:rPr")
            sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(font_size * 2))); rpr.append(sz)
            r.append(rpr)
        r.append(child)
        paragraph._p.append(r)
        return r

    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); b.set(qn("w:dirty"), "true"); r_with(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r_with(it)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r_with(s)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = placeholder; r_with(t)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r_with(e)


def set_pgnum(section, fmt: str | None, start: int | None):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    pg = OxmlElement("w:pgNumType")
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    cols = sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(pg)
    else:
        sectPr.append(pg)


def page_setup(section):
    section.page_width = Inches(8.27)      # A4 210 mm
    section.page_height = Inches(11.69)    # A4 297 mm
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)     # binding gutter
    section.right_margin = Inches(1.0)
    section.gutter = Inches(0)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    # A4 exactly in twips (python-docx rounds inches): 11906 x 16838
    pgSz = section._sectPr.find(qn("w:pgSz"))
    pgSz.set(qn("w:w"), "11906")
    pgSz.set(qn("w:h"), "16838")
    pgMar = section._sectPr.find(qn("w:pgMar"))
    for k, v in {"top": "1440", "bottom": "1440", "left": "1800", "right": "1440"}.items():
        pgMar.set(qn("w:" + k), v)


# ------------------------------------------------------------------ inline markup: **bold**, *italic*, `code`
TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*|_\{.+?\}|\^\{.+?\})")


REF_RESOLVER = None   # set by Report: maps "@[key]" -> "Figure 5.3" / "Table 2.1"


def add_marked_runs(paragraph, text: str, size=None, mono_size=None, base_italic=False, base_bold=False):
    if REF_RESOLVER is not None:
        text = re.sub(r"@\[(\w+)\]", lambda m: REF_RESOLVER(m.group(1)), text)
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True; r.italic = base_italic
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1]); set_rfonts(r._r.get_or_add_rPr(), MONO_FONT)
            r.font.size = Pt(mono_size or ((size or 12) - 1.5)); r.bold = base_bold; r.italic = base_italic
        elif part.startswith("_{") and part.endswith("}"):
            r = paragraph.add_run(part[2:-1]); r.font.subscript = True; r.italic = base_italic; r.bold = base_bold
        elif part.startswith("^{") and part.endswith("}"):
            r = paragraph.add_run(part[2:-1]); r.font.superscript = True; r.italic = base_italic; r.bold = base_bold
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = paragraph.add_run(part[1:-1]); r.italic = True; r.bold = base_bold
        else:
            r = paragraph.add_run(part); r.italic = base_italic; r.bold = base_bold
        if size and not part.startswith("`"):
            r.font.size = Pt(size)
    return paragraph


# ------------------------------------------------------------------ the builder
class Report:
    def __init__(self, template_initial: Path, labels: dict | None = None, chapter_page_breaks: bool = True):
        self.chapter_page_breaks = chapter_page_breaks
        self.doc = Document()
        self.template_initial = Path(template_initial)
        self.labels = labels or {}          # key -> "Figure 5.3" (from the previous pass)
        self.collected: dict[str, str] = {}
        self.chapter = None                 # "1".."7", "A", "B" or None
        self.fig_n = 0
        self.tab_n = 0
        self.code_n = 0
        self.warnings: list[str] = []
        self._setup_styles()
        page_setup(self.doc.sections[0])
        global REF_RESOLVER
        REF_RESOLVER = self.ref

    # ---------------------------------------------------------------- styles
    def _setup_styles(self):
        d = self.doc
        st = d.styles
        defaults = st.element.find(qn("w:docDefaults"))
        set_rfonts(defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")), BODY_FONT)

        normal = st["Normal"]
        set_rfonts(normal.element.get_or_add_rPr(), BODY_FONT)
        normal.font.size = Pt(12)
        pf = normal.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.widow_control = True

        for name, size, align, before, after, caps in [
            ("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER, 0, 18, True),
            ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT, 12, 6, False),
            ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT, 10, 4, False),
        ]:
            s = st[name]
            rpr = s.element.get_or_add_rPr()
            set_rfonts(rpr, BODY_FONT)
            set_color(rpr, "000000")
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.italic = False
            s.font.all_caps = caps
            p = s.paragraph_format
            p.alignment = align
            p.space_before = Pt(before)
            p.space_after = Pt(after)
            p.line_spacing = 1.15
            p.keep_with_next = True
            p.keep_together = True
        st["Heading 1"].paragraph_format.page_break_before = self.chapter_page_breaks
        if not self.chapter_page_breaks:
            st["Heading 1"].paragraph_format.space_before = Pt(18)

        def para_style(name, size, align, italic=False, bold=False, font=BODY_FONT, before=0, after=0,
                       line=1.0, keep_next=False):
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = st["Normal"]
            set_rfonts(s.element.get_or_add_rPr(), font)
            s.font.size = Pt(size)
            s.font.italic = italic
            s.font.bold = bold
            p = s.paragraph_format
            p.alignment = align
            p.space_before = Pt(before)
            p.space_after = Pt(after)
            p.line_spacing = line
            p.keep_with_next = keep_next
            s.quick_style = True
            return s

        para_style("Prelim Title", 16, WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=18, line=1.15)
        st["Prelim Title"].font.all_caps = True
        st["Prelim Title"].paragraph_format.page_break_before = True
        para_style("Prelim Heading", 16, WD_ALIGN_PARAGRAPH.CENTER, bold=True, before=14, after=8, line=1.15,
                   keep_next=True)
        st["Prelim Heading"].font.all_caps = True
        para_style("Figure Caption", 10, WD_ALIGN_PARAGRAPH.CENTER, italic=True, before=4, after=12, line=1.0)
        para_style("Table Caption", 10, WD_ALIGN_PARAGRAPH.CENTER, italic=True, before=8, after=4, line=1.0,
                   keep_next=True)
        para_style("Table Text", 10, WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.0)
        para_style("Code", 9, WD_ALIGN_PARAGRAPH.LEFT, font=MONO_FONT, after=0, line=1.0)
        para_style("Code Label", 10, WD_ALIGN_PARAGRAPH.LEFT, italic=True, before=6, after=2, line=1.0,
                   keep_next=True)
        para_style("Equation", 10, WD_ALIGN_PARAGRAPH.LEFT, font=MONO_FONT, before=4, after=8, line=1.0)
        para_style("Figure Body", 12, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=0, line=1.0, keep_next=True)
        para_style("Note", 10, WD_ALIGN_PARAGRAPH.JUSTIFY, italic=True, before=2, after=6, line=1.0)

        lb = st["List Bullet"]
        set_rfonts(lb.element.get_or_add_rPr(), BODY_FONT)
        lb.font.size = Pt(12)
        lb.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        lb.paragraph_format.line_spacing = 1.5
        lb.paragraph_format.space_after = Pt(3)
        # Word's built-in contents styles are named "toc 1".."toc 3" (style IDs TOC1..TOC3); defining them under
        # those exact names makes Word use these properties when it rebuilds the TOC / LOF / LOT fields.
        for level in (1, 2, 3):
            s = st.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH)
            s.element.set(qn("w:styleId"), f"TOC{level}")
            s.base_style = st["Normal"]
            set_rfonts(s.element.get_or_add_rPr(), BODY_FONT)
            s.font.size = Pt(12 if level == 1 else 11)
            s.font.bold = level == 1
            s.paragraph_format.line_spacing = 1.0
            s.paragraph_format.space_after = Pt(2)
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            s.paragraph_format.left_indent = Inches({1: 0, 2: 0.3, 3: 0.6}[level])
            s.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT,
                                                      WD_TAB_LEADER.DOTS)

    # ---------------------------------------------------------------- numbering helpers
    def set_chapter(self, ch):
        self.chapter, self.fig_n, self.tab_n, self.code_n = ch, 0, 0, 0

    def ref(self, key: str) -> str:
        return self.labels.get(key, "Figure ?" if key.startswith("f_") else "Table ?")

    # ---------------------------------------------------------------- text
    def h1(self, text, chapter=None, new_page=False):
        if chapter is not None:
            self.set_chapter(chapter)
        para = self.doc.add_paragraph(text, style="Heading 1")
        if new_page:
            para.paragraph_format.page_break_before = True
        return para

    def h2(self, text):
        return self.doc.add_paragraph(text, style="Heading 2")

    def h3(self, text):
        return self.doc.add_paragraph(text, style="Heading 3")

    def prelim_title(self, text):
        return self.doc.add_paragraph(text, style="Prelim Title")

    def prelim_heading(self, text):
        return self.doc.add_paragraph(text, style="Prelim Heading")

    def p(self, text, align=None, italic=False, size=None, space_after=None, keep_next=False):
        para = self.doc.add_paragraph(style="Normal")
        add_marked_runs(para, text, size=size, base_italic=italic)
        if align:
            para.alignment = align
        if space_after is not None:
            para.paragraph_format.space_after = Pt(space_after)
        if keep_next:
            para.paragraph_format.keep_with_next = True
        return para

    def bullets(self, items):
        for it in items:
            add_marked_runs(self.doc.add_paragraph(style="List Bullet"), it)

    def note(self, text):
        return add_marked_runs(self.doc.add_paragraph(style="Note"), text, size=10, base_italic=True)

    def eq(self, expr: str, num: str | None = None):
        para = self.doc.add_paragraph(style="Equation")
        para.paragraph_format.left_indent = Inches(0.2)
        para.paragraph_format.right_indent = Inches(0.0)
        para.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN - 0.25), WD_TAB_ALIGNMENT.RIGHT)
        add_marked_runs(para, expr + (f"\t{num}" if num else ""), size=10)
        for r in para.runs:
            if not r.font.name:
                set_rfonts(r._r.get_or_add_rPr(), MONO_FONT)
        para_shading(para)
        return para

    def code(self, lines, label=None, size=9):
        if label:
            if label.startswith("Code excerpt"):
                self.code_n += 1
            add_marked_runs(self.doc.add_paragraph(style="Code Label"), label, size=10, base_italic=True)
        lines = list(lines)
        while lines and not lines[-1].strip():
            lines.pop()
        for i, line in enumerate(lines):
            para = self.doc.add_paragraph(style="Code")
            para.paragraph_format.left_indent = Inches(0.35)
            para.paragraph_format.first_line_indent = Inches(-0.25)   # hanging indent for wrapped lines
            if i < len(lines) - 1:
                para.paragraph_format.keep_with_next = len(lines) <= 25
            run = para.add_run(line.rstrip("\n").replace("\t", "    ") or " ")
            run.font.size = Pt(size)
            para_shading(para)
            para_borders(para)
        spacer = self.doc.add_paragraph(style="Normal")
        spacer.paragraph_format.space_after = Pt(2)
        spacer.paragraph_format.line_spacing = 1.0

    # ---------------------------------------------------------------- figures & tables
    def _label(self, kind, key):
        if kind == "Figure":
            self.fig_n += 1
            n = self.fig_n
        else:
            self.tab_n += 1
            n = self.tab_n
        label = f"{kind} {self.chapter}.{n}"
        if key in self.collected:
            raise ValueError(f"duplicate key {key}")
        self.collected[key] = label
        return label

    def fig(self, key, path, caption, width_in=TEXT_WIDTH_IN, max_h_in=MAX_FIG_HEIGHT_IN):
        path = Path(path)
        if not path.exists():
            self.warnings.append(f"missing figure {path}")
            label = self._label("Figure", key)
            self.p(f"[PLACEHOLDER — {path.name} not found]", align=WD_ALIGN_PARAGRAPH.CENTER)
            cap = self.doc.add_paragraph(style="Figure Caption")
            add_marked_runs(cap, f"{label}: {caption}", size=10, base_italic=True)
            return label
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
        width_in = min(width_in, TEXT_WIDTH_IN, max_h_in * w / h)
        para = self.doc.add_paragraph(style="Figure Body")
        para.add_run().add_picture(str(path), width=Inches(width_in))
        label = self._label("Figure", key)
        cap = self.doc.add_paragraph(style="Figure Caption")
        add_marked_runs(cap, f"{label}: {caption}", size=10, base_italic=True)
        return label

    def table(self, key, caption, header, rows, widths, align=None, font_size=10, bold_first_col=False,
              repeat_header=True, caption_style="Table Caption", keep_together=None):
        """widths in inches; align: list of 'l'/'c'/'r' per column."""
        font_size = 10              # guideline: tables 10 pt
        total = sum(widths)
        if total > TEXT_WIDTH_IN + 0.02:
            raise ValueError(f"table {key} too wide: {total:.2f} in")
        label = None
        if caption is not None:
            label = self._label("Table", key)
            cap = self.doc.add_paragraph(style=caption_style)
            add_marked_runs(cap, f"{label}: {caption}", size=10, base_italic=True)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
        align = align or ["l"] * len(header)
        amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT,
                "j": WD_ALIGN_PARAGRAPH.JUSTIFY}
        for ri, row_vals in enumerate([header] + list(rows)):
            row = t.rows[ri]
            trPr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit"); trPr.append(cant)
            if ri == 0 and repeat_header:
                th = OxmlElement("w:tblHeader"); trPr.append(th)
            for ci, val in enumerate(row_vals):
                cell = row.cells[ci]
                cell.width = Inches(widths[ci])
                para = cell.paragraphs[0]
                para.style = self.doc.styles["Table Text"]
                para.alignment = amap[align[ci]] if ri else WD_ALIGN_PARAGRAPH.CENTER
                parts = str(val).split("\n")
                for pi, part in enumerate(parts):
                    target = para if pi == 0 else cell.add_paragraph(style="Table Text")
                    if pi:
                        target.alignment = para.alignment
                    add_marked_runs(target, part, size=font_size, mono_size=font_size,
                                    base_bold=(ri == 0) or (bold_first_col and ci == 0))
                if ri == 0:
                    shade(cell._tc.get_or_add_tcPr(), HEADER_FILL)
                # keep short tables on one page; always keep the header row with the first data row
                together = (len(rows) <= 12) if keep_together is None else keep_together
                if ri == 0 or (together and ri < len(rows)):
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        for ci, wdt in enumerate(widths):
            for cell in t.columns[ci].cells:
                cell.width = Inches(wdt)
        # explicit grid widths
        grid = t._tbl.tblGrid
        for gc, wdt in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(wdt * 1440)))
        tail = self.doc.add_paragraph(style="Normal")
        tail.paragraph_format.space_after = Pt(4)
        tail.paragraph_format.line_spacing = 1.0
        return label

    def page_break(self):
        from docx.enum.text import WD_BREAK
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def field_paragraph(self, instr, placeholder):
        para = self.doc.add_paragraph(style="Normal")
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_field(para, instr, placeholder)
        return para

    # ---------------------------------------------------------------- sections, headers, footers
    def new_section(self):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        page_setup(sec)
        return sec

    @staticmethod
    def footer_page_number(footer):
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        add_field(para, "PAGE", "1", font_size=11)
        for r in para._p.findall(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr"); r.insert(0, rpr)
            set_rfonts(rpr, BODY_FONT)

    def logo_bytes(self):
        with zipfile.ZipFile(self.template_initial) as z:
            return z.read("word/media/image1.png")

    def college_header(self, header):
        """Reproduce the Initial Pages template header: logo + program line + honours line + spacer."""
        header.is_linked_to_previous = False
        for para in list(header.paragraphs):
            para._p.getparent().remove(para._p)
        p1 = header.add_paragraph(style="Header")
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.add_run().add_picture(io.BytesIO(self.logo_bytes()), width=Inches(TEXT_WIDTH_IN))
        for text, font, size, bold in [
            ("UG Program in Electronics and Computer Science", "Aptos Display", 16, False),
            ("Honours in Emerging Area offered by Department of Artificial Intelligence and Data Science",
             "Aptos", 14, True),
            ("", "Aptos", 14, True),
        ]:
            para = header.add_paragraph(style="Header")
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(8)
            para.paragraph_format.space_after = Pt(4)
            para.paragraph_format.line_spacing = 1.16
            if text:
                r = para.add_run(text)
                rpr = r._r.get_or_add_rPr()
                set_rfonts(rpr, font)
                r.font.size = Pt(size)
                r.bold = bold
                r.font.color.rgb = RGBColor(0x0F, 0x47, 0x61)
            else:
                rpr = para._p.get_or_add_pPr()
                rr = OxmlElement("w:rPr"); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(size * 2))
                rr.append(sz); rpr.append(rr)

    # ---------------------------------------------------------------- initial pages copied from the template
    def copy_initial_pages(self, fill: dict):
        tpl = Document(str(self.template_initial))
        # bring over any style the template pages reference that this document lacks
        my_ids = {s.get(qn("w:styleId")) for s in self.doc.styles.element.findall(qn("w:style"))}
        tpl_styles = {s.get(qn("w:styleId")): s for s in tpl.styles.element.findall(qn("w:style"))}
        body_children = [c for c in tpl.element.body if c.tag != qn("w:sectPr")]
        referenced = set()
        for c in body_children:
            for tag in ("w:pStyle", "w:rStyle", "w:tblStyle"):
                for e in c.iter(qn(tag)):
                    referenced.add(e.get(qn("w:val")))
        for sid in sorted(referenced - my_ids):
            if sid in tpl_styles:
                self.doc.styles.element.append(copy.deepcopy(tpl_styles[sid]))

        def text_of(el):
            return "".join(t.text or "" for t in el.iter(qn("w:t")))

        starts = {"Certificate": None, "Approval for Project Report for T. Y. Semester V": None, "Declaration": None}
        for i, c in enumerate(body_children):
            tx = text_of(c).strip()
            if c.tag == qn("w:p") and tx in starts and starts[tx] is None:
                starts[tx] = i
        i_cert, i_appr, i_decl = starts.values()
        pages = {"cover": body_children[:i_cert], "certificate": body_children[i_cert:i_appr],
                 "approval": body_children[i_appr:i_decl], "declaration": body_children[i_decl:]}

        def is_empty_p(el):
            return el.tag == qn("w:p") and not text_of(el).strip() and not list(el.iter(qn("w:drawing")))

        for k, els in pages.items():
            while els and is_empty_p(els[-1]):
                els.pop()
            pages[k] = els

        body = self.doc.element.body
        sectPr = body.find(qn("w:sectPr"))
        order = [("cover", "Outer cover"), ("cover", "Inner title page"), ("certificate", "Certificate"),
                 ("approval", "Approval"), ("declaration", "Declaration")]
        if not fill.get("inner_title", True):
            order = [o for o in order if o[1] != "Inner title page"]
        for n, (key, _name) in enumerate(order):
            new_els = [copy.deepcopy(e) for e in pages[key]]
            for e in new_els:
                self._normalise_template_element(e)
            self._fill(key, new_els, fill)
            if n > 0:
                first_p = next(e for e in new_els if e.tag == qn("w:p"))
                Paragraph(first_p, self.doc._body).paragraph_format.page_break_before = True
            for e in new_els:
                sectPr.addprevious(e)

    def _normalise_template_element(self, el):
        """Re-create the template's inherited defaults explicitly (it relied on Aptos, left alignment,
        8 pt after / 1.16 line spacing), because this document's Normal style is Times New Roman, justified."""
        for p in el.iter(qn("w:p")):
            pPr = p.find(qn("w:pPr"))
            style = pPr.find(qn("w:pStyle")).get(qn("w:val")) if pPr is not None and pPr.find(qn("w:pStyle")) is not None else "Normal"
            par = Paragraph(p, self.doc._body)
            if style == "Normal":
                sp = p.find(qn("w:pPr"))
                sp_el = sp.find(qn("w:spacing")) if sp is not None else None
                if sp_el is None or sp_el.get(qn("w:after")) is None:
                    par.paragraph_format.space_after = Pt(8)
                if sp_el is None or sp_el.get(qn("w:line")) is None:
                    par.paragraph_format.line_spacing = 278 / 240
                if sp_el is None or sp_el.get(qn("w:before")) is None:
                    par.paragraph_format.space_before = Pt(0)
                if pPr is None or pPr.find(qn("w:jc")) is None:
                    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in list(p.iter(qn("w:r"))):
                rpr = r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr"); r.insert(0, rpr)
                rf = rpr.find(qn("w:rFonts"))
                if rf is None or rf.get(qn("w:ascii")) is None:
                    major = rf is not None and (rf.get(qn("w:asciiTheme")) or "").startswith("major")
                    set_rfonts(rpr, "Aptos Display" if major else "Aptos")
            ppr = p.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:rPr")) is not None:
                mrpr = ppr.find(qn("w:rPr"))
                rf = mrpr.find(qn("w:rFonts"))
                if rf is None or rf.get(qn("w:ascii")) is None:
                    set_rfonts(mrpr, "Aptos")

    @staticmethod
    def _set_run_text(r, text):
        ts = r.findall(qn("w:t"))
        if not ts:
            t = OxmlElement("w:t"); r.append(t); ts = [t]
        ts[0].text = text
        ts[0].set(qn("xml:space"), "preserve")
        for extra in ts[1:]:
            r.remove(extra)

    def _append_run(self, p, text):
        runs = p.findall(qn("w:r"))
        new = copy.deepcopy(runs[-1])
        for child in list(new):
            if child.tag != qn("w:rPr"):
                new.remove(child)
        t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); new.append(t)
        runs[-1].addnext(new)

    def _style_run(self, r, font, size_pt, bold):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr"); r.insert(0, rpr)
        set_rfonts(rpr, font)
        for tag in ("w:b", "w:bCs", "w:sz", "w:szCs"):
            for e in rpr.findall(qn(tag)):
                rpr.remove(e)
        if bold:
            b = OxmlElement("w:b"); rpr.append(b)
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size_pt * 2))); rpr.append(sz)

    @staticmethod
    def _set_tbl_widths(tbl, widths_in):
        tw = [str(int(w * 1440)) for w in widths_in]
        grid = tbl.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), tw):
            gc.set(qn("w:w"), w)
        tblPr = tbl.find(qn("w:tblPr"))
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
        tblW.set(qn("w:w"), str(sum(int(w) for w in tw))); tblW.set(qn("w:type"), "dxa")
        if tblPr.find(qn("w:tblLayout")) is None:
            lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
            look = tblPr.find(qn("w:tblLook"))
            (look.addprevious(lay) if look is not None else tblPr.append(lay))
        for tr in tbl.findall(qn("w:tr")):
            for tc, w in zip(tr.findall(qn("w:tc")), tw):
                tcPr = tc.find(qn("w:tcPr"))
                if tcPr is None:
                    tcPr = OxmlElement("w:tcPr"); tc.insert(0, tcPr)
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = OxmlElement("w:tcW"); tcPr.insert(0, tcW)
                tcW.set(qn("w:w"), w); tcW.set(qn("w:type"), "dxa")

    def _fill(self, key, els, fill):
        def ptext(el):
            return "".join(t.text or "" for t in el.iter(qn("w:t")))

        title = fill["title"]
        if key == "cover":
            for e in els:
                if e.tag != qn("w:p"):
                    continue
                tx = ptext(e).strip()
                if re.fullmatch(r"-{10,}", tx):
                    r = e.findall(qn("w:r"))[0]
                    self._set_run_text(r, title)
                    self._style_run(r, "Aptos", 16, True)
                elif tx == "Name:":
                    self._append_run(e, " " + fill["name"])
                elif tx == "PRN:":
                    self._append_run(e, " " + fill["prn"])
                elif tx == "Class and Division:":
                    self._append_run(e, " " + fill["class_div"])
                elif tx == "Guide:":
                    self._append_run(e, " " + fill["guide"])
        elif key == "certificate":
            for e in els:
                if e.tag == qn("w:tbl"):
                    self._set_tbl_widths(e, [2.3, 0.9, 1.7])
            done_title = False
            for e in els:
                if e.tag == qn("w:p") and re.fullmatch(r"-{40,}", ptext(e).strip()) and not done_title:
                    r = e.findall(qn("w:r"))[0]
                    self._set_run_text(r, title)
                    self._style_run(r, "Times New Roman", 14, True)
                    done_title = True
                elif e.tag == qn("w:tbl"):
                    for t in e.iter(qn("w:t")):
                        if re.fullmatch(r"\s*-{10,}\s*", t.text or ""):
                            t.text = fill["name"]
                        elif re.fullmatch(r"\s*-{5,9}\s*", t.text or ""):
                            t.text = fill["roll"]
        elif key == "approval":
            for e in els:
                if e.tag != qn("w:p"):
                    continue
                for t in e.iter(qn("w:t")):
                    if t.text and re.search(r"_{20,}", t.text) and "“" in t.text:
                        t.text = re.sub(r"_{20,}", title, t.text)
                    elif t.text and t.text.startswith("Student Name") and "____" in t.text:
                        t.text = re.sub(r"_{5,}", fill["name"], t.text)
                    elif t.text and re.fullmatch(r"1\._{20,}", t.text.strip()):
                        t.text = "1." + "_" * 22   # examiner signature line shortened to fit the A4 text width
        elif key == "declaration":
            for e in els:
                if e.tag == qn("w:tbl"):
                    self._set_tbl_widths(e, [2.1, 0.8, 1.9, 1.2])
                    trs = e.findall(qn("w:tr"))
                    new_tr = copy.deepcopy(trs[0])
                    values = [fill["name"], "T. Y.", fill["roll"], ""]
                    for tc, val in zip(new_tr.findall(qn("w:tc")), values):
                        ps = tc.findall(qn("w:p"))
                        for extra in ps[1:]:
                            tc.remove(extra)
                        runs = ps[0].findall(qn("w:r"))
                        for extra in runs[1:]:
                            ps[0].remove(extra)
                        if runs:
                            self._set_run_text(runs[0], val)
                            rpr = runs[0].find(qn("w:rPr"))
                            if rpr is not None:
                                for b in rpr.findall(qn("w:b")) + rpr.findall(qn("w:bCs")):
                                    rpr.remove(b)
                    trs[-1].addnext(new_tr)
